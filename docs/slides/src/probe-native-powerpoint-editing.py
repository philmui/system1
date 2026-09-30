#!/usr/bin/env python3
"""Exercise ordinary native editing through python-pptx on an isolated copy.

Requires python-pptx. This is a package/model round trip, not an Office render.
The delivered input is never overwritten and temporary edited slides are not
published as a tutorial revision.
"""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import tempfile
from zipfile import ZipFile

from pptx import Presentation, __version__ as pptx_version
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)


def native_text(shape):
    return shape._element.find(qn("p:txBody")) is not None


def named(pres, slide, fragment):
    matches = [shape for shape in walk(pres.slides[slide - 1].shapes) if fragment in shape.name]
    if len(matches) != 1:
        raise AssertionError(f"Expected exactly one component matching {fragment!r}: {len(matches)}")
    return matches[0]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()
    before = digest(args.pptx)
    pres = Presentation(args.pptx)
    all_shapes = [shape for slide in pres.slides for shape in walk(slide.shapes)]
    counts = Counter()
    for shape in all_shapes:
        counts["groups" if shape.shape_type == MSO_SHAPE_TYPE.GROUP else
               "textComponents" if native_text(shape) else "drawingComponents"] += 1
    assert len(pres.slides) == 36
    assert counts["textComponents"] == 722
    assert counts["drawingComponents"] == 980
    assert all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in all_shapes)

    title = named(pres, 1, "[p01-text-0013]")
    assert native_text(title)
    old_title = title.text
    title.text_frame.paragraphs[0].runs[0].text = "Customized enterprise"
    new_title = title.text
    assert old_title != new_title and "Customized enterprise" in new_title

    bar = named(pres, 24, "[p24-e0025]")
    old_bar_width = bar.width
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string("F6B7C3")
    bar.width = old_bar_width - 127000

    # A semantic component inside the library, not the entire slide frame.
    groups = [shape for shape in walk(pres.slides[35].shapes)
              if shape.shape_type == MSO_SHAPE_TYPE.GROUP and shape.name.startswith("whole-kit ·")]
    if not groups:
        raise AssertionError("No editable whole-kit illustration group found in the component library")
    group = groups[-1]
    group_name = group.name
    old_left = group.left
    group.left = old_left + 63500

    with tempfile.TemporaryDirectory(prefix="system1-native-edit-probe-") as folder:
        edited = Path(folder) / "edited-copy.pptx"
        pres.save(edited)
        again = Presentation(edited)
        assert named(again, 1, "[p01-text-0013]").text == new_title
        edited_bar = named(again, 24, "[p24-e0025]")
        assert str(edited_bar.fill.fore_color.rgb) == "F6B7C3"
        assert edited_bar.width == old_bar_width - 127000
        edited_group = [shape for shape in walk(again.slides[35].shapes) if shape.name == group_name]
        assert len(edited_group) == 1 and edited_group[0].left == old_left + 63500
        with ZipFile(edited) as archive:
            assert archive.testzip() is None
        edited_hash = digest(edited)
    assert digest(args.pptx) == before
    report = {
        "passed": True,
        "method": f"python-pptx {pptx_version}: read native object model, edit isolated copy, save and reopen",
        "input": str(args.pptx), "inputSha256": before,
        "inputUnchanged": True, "slideCount": len(pres.slides),
        "componentsReadThroughNativeAPI": dict(counts),
        "editsVerifiedAfterReopen": [
            {"kind": "text", "slide": 1, "component": title.name, "before": old_title, "after": new_title},
            {"kind": "fill", "slide": 24, "component": bar.name, "afterRGB": "F6B7C3"},
            {"kind": "resize", "slide": 24, "component": bar.name, "widthChangeEMU": -127000},
            {"kind": "moveGroup", "slide": 36, "component": group_name, "xChangeEMU": 63500},
        ],
        "temporaryCopySha256": edited_hash,
        "temporaryCopyRemoved": True,
        "nativeOfficeRendered": False,
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: report[k] for k in ("passed", "method", "componentsReadThroughNativeAPI", "inputUnchanged")}))


if __name__ == "__main__":
    main()
